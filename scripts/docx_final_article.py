"""Editable Word text and tables with the original article illustrations."""
import json,re
from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
doc=Document();section=doc.sections[0]
section.page_width=Mm(210);section.page_height=Mm(297)
section.top_margin=Mm(17);section.bottom_margin=Mm(17)
section.left_margin=Mm(18);section.right_margin=Mm(18)
for name,size in [('Normal',11),('Title',18),('Heading 1',13),('Heading 2',11.5),('Caption',9)]:
    style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_after=Pt(6)
    style.paragraph_format.line_spacing=1.08
    if name!='Normal':style.paragraph_format.keep_with_next=True
doc.styles['Title'].font.bold=True
doc.styles['Normal'].paragraph_format.widow_control=True
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def inline(p,text):
    for part in re.split(r'(\[[^\]]+\]\([^)]+\))',text):
        m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part)
        if m:
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(m[2],RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');t=OxmlElement('w:t');t.text=m[1];run.append(t);link.append(run);p._p.append(link)
        else:p.add_run(part)

def picture(path,width,description):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    shape=p.add_run().add_picture(str(path),width=Pt(width));shape._inline.docPr.set('descr',description)
    return p

equations=json.loads((ROOT/'article/figures/final20261007/equations.json').read_text())
lines=(ROOT/'article/ARTICLE_FINAL_20261007.md').read_text(encoding='utf-8').splitlines();i=0
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line:continue
    if line=='<!--pagebreak-->':doc.add_page_break();continue
    if line.startswith('$$'):
        formula=line[2:-2];path=Path(equations[formula])
        with Image.open(path) as im:w,h=im.size
        picture(path,min(485,w*72/250),formula);continue
    if line.startswith('!['):
        caption,path=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line).groups()
        picture(ROOT/'article'/path,485,caption)
        p=doc.add_paragraph(style='Caption');inline(p,caption);p.paragraph_format.keep_with_next=False
        continue
    if line.startswith('|'):
        rows=[line]
        while i<len(lines) and lines[i].startswith('|'):rows.append(lines[i]);i+=1
        data=[]
        for row in rows:
            cells=[c.strip() for c in row.strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+',c) for c in cells):data.append(cells)
        n=len(data[0]);widths={3:[100,37,37],4:[54,40,40,40],5:[47,26,34,33,34]}.get(n,[174/n]*n)
        table=doc.add_table(rows=0,cols=n);table.autofit=False
        for col,width in zip(table.columns,widths):col.width=Mm(width)
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
        table._tbl.tblPr.append(borders)
        for ri,row in enumerate(data):
            tr=table.add_row();pr=tr._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            if ri==0:pr.append(OxmlElement('w:tblHeader'))
            for ci,(cell,text) in enumerate(zip(tr.cells,row)):
                cell.width=Mm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp=cell._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+side);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');margin.append(el)
                cp.append(margin)
                if ri==0:
                    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E7EEF1');cp.append(fill)
                p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                inline(p,text)
                for r in p.runs:r.font.size=Pt(9);r.bold=ri==0
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
        continue
    for prefix,style in [('### ','Heading 2'),('## ','Heading 1'),('# ','Title')]:
        if line.startswith(prefix):
            text=line[len(prefix):]
            text=re.sub(r'^(\d+)\. ',r'\1 ',text)
            p=doc.add_paragraph(style=style);inline(p,text);break
    else:
        p=doc.add_paragraph();inline(p,line)
doc.core_properties.title=lines[0][2:];doc.core_properties.author=''
out=ROOT/'output/SMF_Load_Veracity_Final_20261007.docx';doc.save(out)
print(out)
